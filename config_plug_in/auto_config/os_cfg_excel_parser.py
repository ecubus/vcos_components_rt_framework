# This Python file uses the following encoding: utf-8
############################################################################
# Copyright (c) 2025 Li Auto Inc. and its affiliates
# Licensed under the Apache License, Version 2.0(the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
############################################################################
import os
import re
import copy
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill, Border, Font, Side

from autosar44 import autosar44
from basic_func_module.autosar_utils import utils
from basic_func_module.autosar_utils.autosar_util import AutosarUtil
from basic_func_module.base_logger import logger
from func_module.module_management.bsw_editor.bsw_editor import ECUC_PARAM_VALUE_LIST, ECUC_NUMERICAL_PARAM_VALUE_LIST, \
    ECUC_TEXTUAL_PARAM_VALUE_LIST, ECUC_ADD_INFO_PARAM_VALUE_LIST
from func_module.module_management.bsw_editor.bsw_editor import ECUC_REFERENCE_VALUE_LIST, ECUC_IREFERENCE_VALUE_LIST, ECUC_REF_VALUE_LIST

CUR_FILE_PATH = os.path.dirname(os.path.realpath(__file__))


class OsCfgExcelImporter:
    def __init__(self, data_management) -> None:
        self.data_management = data_management
        self.autosar_bsw_class = data_management.autosar_bsw_class
        self.autosar_bswmd_class = data_management.autosar_bswmd_class

        self.format_info = {}           # excel格式信息
        self.workbook = None            # excel文件
        self.os_excel_path = None       # excel文件路径
        self.excel_version = None       # excel文件版本号

        self.app_alarm_mapping = {}     # app的alarm ref
        self.app_isr_mapping = {}       # app的isr ref
        self.task_event_mapping = {}    # task的event ref

    def import_excel(self, os_excel_path) -> bool:
        """导入os自动配置"""
        if os.path.exists(os_excel_path) is False:
            logger.ui_error(f"{os_excel_path} does not exist.")
            return False
        self.os_excel_path = os_excel_path
        self.workbook = load_workbook(os_excel_path)

        if self.validate() is False:
            return False
        self.parse_excel_and_auto_cfg()
        return True

    def validate(self) -> bool:
        """校验excel"""
        if self.validate_format() is False:
            return False
        if self.validate_version() is False:
            return False
        if self.validate_module() is False:
            return False
        if self.validate_header() is False:
            return False
        return True

    def validate_format(self) -> bool:
        """检查格式信息文件"""
        excel_format_path = os.path.join(CUR_FILE_PATH, "excel_format.json")
        if not os.path.exists(excel_format_path):
            logger.ui_error(f"'{excel_format_path}' does not exist.")
            return False
        with open(excel_format_path, "r", encoding="utf-8") as f:
            self.format_info = json.load(f)
        return True

    def validate_version(self) -> bool:
        """检查版本"""
        if "Version" not in self.workbook.sheetnames:
            logger.ui_error(f"Version sheet is not in {self.os_excel_path}.")
            return False
        version_sheet = self.workbook["Version"]
        # 是否有版本字段
        if version_sheet.cell(row=1, column=1).value != "Version":
            logger.ui_error(f"The format of 'Version' sheet is wrong, 'Version' should be in 1 row and 1 column.")
            return False
        # 版本是否有效
        self.excel_version = None
        for row in range(2, version_sheet.max_row + 1):
            self.excel_version = str(version_sheet.cell(row=row, column=1).value)
            if self.excel_version is not None:
                break
        if self.excel_version is None or self.excel_version not in self.format_info:
            logger.ui_error(f"The version number '{self.excel_version}' in Version sheet is not supported({self.os_excel_path}).")
            return False
        return True

    def validate_module(self) -> bool:
        """检查模块是否支持和存在"""
        for module in self.format_info[self.excel_version]["releated_module_list"]:
            if module not in self.autosar_bsw_class:
                logger.ui_error(f"{module} module exists in 'releated_module_list',but it does not exist in current project.")
                return False
            if module not in self.autosar_bswmd_class:
                logger.ui_error(f"{module} module exists in 'releated_module_list',but {module}_bswmd.arxml does not exist in current project.")
                return False
        return True

    def validate_header(self) -> bool:
        """检查excel文件的表头"""
        for sheet_name, sheet_info in self.format_info[self.excel_version]["sheet"].items():
            # sheet是否存在
            if sheet_name not in self.workbook.sheetnames:
                logger.ui_error(f"'{sheet_name}' sheet is not in {self.os_excel_path}.")
                return False
            sheet = self.workbook[sheet_name]
            for header, item in sheet_info["header"].items():
                # 表头是否匹配
                coor = item[0]
                if coor is not None and sheet.cell(row=coor[0], column=coor[2]).value != header:
                    logger.ui_error(f"The format of '{sheet_name}' sheet is wrong, '{header}' should be in {coor[0]} row and {coor[2]} column.")
                    return False
                # mod arxml格式是否匹配
                mod_path = item[1]
                if mod_path is None:
                    continue
                module = mod_path.split("/")[3]
                if module not in self.autosar_bswmd_class:
                    logger.ui_error(f"{module}_bswmd.arxml does not exist in current project.")
                    return False
                if AutosarUtil.find(self.autosar_bswmd_class[module], mod_path) is None:
                    logger.ui_error(f"{header}'s mod_path {mod_path} does not exist in {module}_bswmd.arxml.")
                    return False
        return True

    def parse_excel_and_auto_cfg(self):
        """解析excel并自动配置"""
        self.parse_ecuc_core()
        self.parse_ecuc_partition()
        self.parse_isr()
        self.parse_event()
        self.parse_task()
        self.parse_app()

    def get_sheet_info_by_header(self, sheet, header_column, name_key):
        info_dict = {}
        cur_name = None
        for row in range(2, sheet.max_row + 1):
            name = sheet.cell(row=row, column=header_column[name_key]).value
            if name is not None:
                cur_name = name
            if cur_name is None:
                continue
            if cur_name not in info_dict:
                info_dict[cur_name] = {}
                info_dict[cur_name].update({key: copy.deepcopy(None) for key in header_column.get("parameter", {})})
                info_dict[cur_name].update({key: copy.deepcopy([]) for key in header_column.get("reference", {})})
            for param, col in header_column.get("parameter", {}).items():
                param_value = sheet.cell(row=row, column=col).value
                if param_value is not None:
                    info_dict[cur_name][param] = param_value
            for param, col in header_column.get("reference", {}).items():
                param_value = sheet.cell(row=row, column=col).value
                if param_value is not None:
                    _ = re.split(r"[、\n\r]+", param_value)
                    info_dict[cur_name][param] += [item.strip() for item in _ if item.strip() != ""]

        return info_dict

    def init_empty_dict(self, sheet_name, module_name, name_key):
        """按表头初始化空dict"""
        header = self.format_info[self.excel_version]["sheet"][sheet_name]["header"]
        bswmd_obj = self.autosar_bswmd_class[module_name]
        base_dict = {}
        for key, item in header.items():
            if key == name_key:
                continue
            mod_path = item[1]
            if mod_path is None:
                base_dict[key] = None
                continue
            mod_obj = bswmd_obj.find(mod_path)
            if mod_obj is None:
                base_dict[key] = None
                continue
            tag = mod_obj.original_tagname_
            if tag in ECUC_PARAM_VALUE_LIST:
                base_dict[key] = None
            elif tag in ECUC_REF_VALUE_LIST:
                if mod_obj.UPPER_MULTIPLICITY is not None and mod_obj.UPPER_MULTIPLICITY.valueOf_ == "1":
                    base_dict[key] = None
                else:
                    base_dict[key] = []
            else:
                base_dict[key] = None
        return base_dict

    def get_sheet_info(self, sheet_name, module_name, name_key):
        """获取sheet信息"""
        sheet = self.workbook[sheet_name]
        header = self.format_info[self.excel_version]["sheet"][sheet_name]["header"]
        base_dict = self.init_empty_dict(sheet_name, module_name, name_key)

        info_dict = {}
        cur_name = None
        for row in range(2, sheet.max_row + 1):
            name = sheet.cell(row=row, column=header[name_key][0][2]).value
            if name is not None:
                cur_name = name
            if cur_name is None:
                continue
            if cur_name not in info_dict:
                info_dict[cur_name] = copy.deepcopy(base_dict)
            for key, item in header.items():
                if key == name_key:
                    continue
                coor = item[0]
                if coor is None:
                    continue
                value = sheet.cell(row=row, column=coor[2]).value
                if value is not None:
                    if isinstance(info_dict[cur_name][key], list) is True:
                        _ = re.split(r"[、\n\r]+", value)
                        info_dict[cur_name][key] += [v.strip() for v in _ if v.strip() != ""]
                    else:
                        info_dict[cur_name][key] = value
        return info_dict

    def parse_ecuc_core(self):
        """解析EcucCoreDefinition并自动配置"""
        # 解析EcucCoreDefinition
        ecuc_core_info_dict = self.get_sheet_info("EcucCoreDefinition", "EcuC", "EcucCoreDefinition(s)")

        ecuc_ecuc = self.autosar_bsw_class["EcuC"]
        ecuc_bswmd = self.autosar_bswmd_class["EcuC"]
        os_ecuc = self.autosar_bsw_class["Os"]
        self.create_ecuc_core(ecuc_ecuc, ecuc_bswmd, os_ecuc, ecuc_core_info_dict)

    def create_ecuc_core(self, ecuc_ecuc, ecuc_bswmd, os_ecuc, ecuc_core_info_dict):
        """配置EcucHardware下的EcucCore"""
        # 移除原有EcucHardware
        OsCfgExecutor.remove_containers(ecuc_ecuc, "/AUTOSAR/EcucDefs/EcuC/EcucHardware")
        # 创建EcucHardware
        ecuc_container = ecuc_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        ecuc_hardware_bswmd_obj = ecuc_bswmd.find("/AUTOSAR/EcucDefs/EcuC/EcucHardware")
        ecuc_hardware_container = AutosarUtil.create_container_value("EcucHardware", ecuc_hardware_bswmd_obj, ecuc_container, "EcucHardware")
        # 获取芯片名称
        chip_name = "chip_name"
        for _ in os_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0].CONTAINERS.ECUC_CONTAINER_VALUE:
            if getattr(_, "DEFINITION_REF", None) is None or getattr(_.DEFINITION_REF, "valueOf_", None) is None:
                continue
            if _.DEFINITION_REF.valueOf_ != "/AUTOSAR/EcucDefs/Os/OsPublishedInformation":
                continue
            if getattr(_, "SHORT_NAME", None) is None or getattr(_.SHORT_NAME, "valueOf_", None) is None:
                continue
            chip_name = _.SHORT_NAME.valueOf_
        if chip_name == "chip_name":
            logger.ui_warning(f"OsPublishedInformation(s) in Os module should have one sub container.")
        # 创建EcucCoreDefinition
        ecuc_core_bswmd_obj = ecuc_bswmd.find("/AUTOSAR/EcucDefs/EcuC/EcucHardware/EcucCoreDefinition")
        for core_name, item in ecuc_core_info_dict.items():
            ecuc_core_container = AutosarUtil.create_container_value(str(core_name), ecuc_core_bswmd_obj, ecuc_hardware_container, str(core_name))
            for param in ecuc_core_container.PARAMETER_VALUES.ECUC_NUMERICAL_PARAM_VALUE:
                if param.DEFINITION_REF.valueOf_ == "/AUTOSAR/EcucDefs/EcuC/EcucHardware/EcucCoreDefinition/EcucCoreId":
                    OsCfgExecutor.set_value(param, item["EcucCoreId"])
                    break
            OsCfgExecutor.modify_single_reference(item["EcucCoreHwRef"], ecuc_core_container, "/AUTOSAR/EcucDefs/EcuC/EcucHardware/EcucCoreDefinition/EcucCoreHwRef", f"/ActiveEcuC/Os/{chip_name}/OsDerivativeInformation/")

    def parse_ecuc_partition(self):
        """解析EcucPartition并自动配置"""
        # 解析EcucPartition
        ecuc_partition_info_dict = self.get_sheet_info("EcucPartition", "EcuC", "EcucPartition(s)")

        ecuc_ecuc = self.autosar_bsw_class["EcuC"]
        ecuc_bswmd = self.autosar_bswmd_class["EcuC"]
        self.create_ecuc_partition(ecuc_ecuc, ecuc_bswmd, ecuc_partition_info_dict)

    def create_ecuc_partition(self, ecuc_ecuc, ecuc_bswmd, ecuc_partition_info_dict):
        """配置EcucPartitionCollection下的EcucPartition"""
        # 移除原有EcucPartitionCollection
        OsCfgExecutor.remove_containers(ecuc_ecuc, "/AUTOSAR/EcucDefs/EcuC/EcucPartitionCollection")
        # 创建EcucPartitionCollection
        ecuc_container = ecuc_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        ecuc_collection_bswmd_obj = ecuc_bswmd.find("/AUTOSAR/EcucDefs/EcuC/EcucPartitionCollection")
        ecuc_collection_container = AutosarUtil.create_container_value("EcucPartitionCollection", ecuc_collection_bswmd_obj, ecuc_container, "EcucPartitionCollection")
        # 创建EcucPartition
        OsCfgExecutor.add_containers(self.format_info, self.excel_version, "EcucPartition", ecuc_bswmd, ecuc_partition_info_dict, ecuc_collection_container)

    def parse_isr(self):
        """解析OsIsr并自动配置"""
        # 解析OsIsr
        isr_info_dict = self.get_sheet_info("OsIsr", "Os", "OsIsr(s)")
        # isr所属Application
        self.app_isr_mapping = {}
        for isr_name, item in isr_info_dict.items():
            if item["所属Application"] is not None:
                if item["所属Application"] not in self.app_isr_mapping:
                    self.app_isr_mapping[item["所属Application"]] = [isr_name]
                else:
                    self.app_isr_mapping[item["所属Application"]].append(isr_name)

        os_bswmd = self.autosar_bswmd_class["Os"]
        os_ecuc = self.autosar_bsw_class["Os"]
        self.create_os_isr(os_ecuc, os_bswmd, isr_info_dict)

    def create_os_isr(self, os_ecuc, os_bswmd, isr_info_dict):
        """配置OsIsr(s)下的isr"""
        # 移除原有isr
        OsCfgExecutor.remove_containers(os_ecuc, "/AUTOSAR/EcucDefs/Os/OsIsr", True)
        os_container = os_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        # 创建isr
        OsCfgExecutor.add_containers(self.format_info, self.excel_version, "OsIsr", os_bswmd, isr_info_dict, os_container)

    def parse_event(self):
        """解析OsEvent并自动配置"""
        event_info_dict = self.get_sheet_info("OsEvent", "Os", "OsEvent(s)")
        # event所属Task
        self.task_event_mapping = {}
        for event_name, item in event_info_dict.items():
            if item["所属Task"] is not None:
                if item["所属Task"] not in self.task_event_mapping:
                    self.task_event_mapping[item["所属Task"]] = [event_name]
                else:
                    self.task_event_mapping[item["所属Task"]].append(event_name)

        os_bswmd = self.autosar_bswmd_class["Os"]
        os_ecuc = self.autosar_bsw_class["Os"]
        self.create_os_event(os_ecuc, os_bswmd, event_info_dict)

    def create_os_event(self, os_ecuc, os_bswmd, event_info_dict):
        """配置OsEvent(s)下的event"""
        # 移除原有event
        OsCfgExecutor.remove_containers(os_ecuc, "/AUTOSAR/EcucDefs/Os/OsEvent", True)
        os_container = os_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        # 创建event
        OsCfgExecutor.add_containers(self.format_info, self.excel_version, "OsEvent", os_bswmd, event_info_dict, os_container)

    def parse_task(self):
        """解析OsTask并自动配置"""
        # 解析OsTask
        task_info_dict = self.get_sheet_info("OsTask", "Os", "OsTask(s)")
        if "Default" not in task_info_dict:
            task_info_dict["Default"] = self.init_empty_dict("OsTask", "Os", "OsTask(s)")

        # 解析APP_TASK
        app_task_sheet = self.workbook["APP_TASK"]
        app_task_header_column = {
            "OsTasks": 2,
            "parameter":
            {
                "Task Type": 3,
                "Task Priority": 4,
                "Task Stack Size [Byte]": 5,
            },
            "reference":
            {
                "Task Event": 6
            }
        }
        temp_info_dict = self.get_sheet_info_by_header(app_task_sheet, app_task_header_column, "OsTasks")

        # 优先以APP_TASK为准
        task_header_column = {
            "OsTask(s)": 1,
            "parameter":
            {
                "OsTaskUsesFpu": 2,
                "OsTaskSchedule": 3,
                "OsTaskType": 4,
                "OsTaskActivation": 5,
                "OsTaskPriority": 6,
                "OsTaskStackSize": 7,
            },
            "reference":
            {
                "OsTaskEventRef": 8,
                "OsTaskAccessingApplication": 9,
            }
        }
        header_mapping = {
            "Task Type": "OsTaskType",
            "Task Priority": "OsTaskPriority",
            "Task Stack Size [Byte]": "OsTaskStackSize",
            "Task Event": "OsTaskEventRef",
        }
        for task_name in temp_info_dict.keys():
            if task_name not in task_info_dict:
                task_info_dict[task_name] = {}
                for k in temp_info_dict[task_name].keys():
                    task_info_dict[task_name][header_mapping[k]] = temp_info_dict[task_name][k]
                # 用默认值补充缺少的key
                for k in task_header_column["parameter"].keys():
                    if k not in task_info_dict[task_name]:
                        task_info_dict[task_name][k] = task_info_dict["Default"][k]
                for k in task_header_column["reference"].keys():
                    if k not in task_info_dict[task_name]:
                        task_info_dict[task_name][k] = task_info_dict["Default"][k]
            else:
                # 覆盖特定key
                for k in temp_info_dict[task_name].keys():
                    task_info_dict[task_name][header_mapping[k]] = temp_info_dict[task_name][k]

        for task_name in task_info_dict.keys():
            task_info_dict[task_name]["OsTaskEventRef"] += self.task_event_mapping.get(task_name, [])
            task_info_dict[task_name]["OsTaskEventRef"] = OsCfgExecutor.remove_list_duplicates(task_info_dict[task_name]["OsTaskEventRef"])

        os_bswmd = self.autosar_bswmd_class["Os"]
        os_ecuc = self.autosar_bsw_class["Os"]
        self.create_os_task(os_ecuc, os_bswmd, task_info_dict)

    def create_os_task(self, os_ecuc, os_bswmd, task_info_dict):
        """配置OsTask(s)下的task"""
        # 移除原有task
        OsCfgExecutor.remove_containers(os_ecuc, "/AUTOSAR/EcucDefs/Os/OsTask", True)
        os_container = os_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        # 创建task
        if "Default" in task_info_dict:
            del task_info_dict["Default"]
        OsCfgExecutor.add_containers(self.format_info, self.excel_version, "OsTask", os_bswmd, task_info_dict, os_container)

        autostart_mod_path = self.format_info[self.excel_version]["sheet"]["OsTask"]["header"]["OsTaskAutostart"][1]
        autostart_suffix = self.format_info[self.excel_version]["sheet"]["OsTask"]["header"]["OsTaskAutostart"][2]
        autostart_mod_obj = os_bswmd.find(autostart_mod_path)
        for container in os_container.CONTAINERS.ECUC_CONTAINER_VALUE:
            if container.DEFINITION_REF.valueOf_ == "/AUTOSAR/EcucDefs/Os/OsTask":
                task_info = task_info_dict[container.SHORT_NAME.valueOf_]
                if "OsTaskAutostart" in task_info and task_info["OsTaskAutostart"]:
                    autostart_con = AutosarUtil.create_container_value("OsTaskAutostart", autostart_mod_obj, container, "OsTaskAutostart")
                    OsCfgExecutor.modify_reference_list(task_info["OsTaskAutostart"], autostart_con, f"{autostart_mod_path}/OsTaskAppModeRef", autostart_suffix)

    def parse_app(self):
        """解析OsApplication并自动配置"""
        # 解析OsApplication
        app_info_dict = self.get_sheet_info("OsApplication", "Os", "OsApplication(s)")
        if "Default" not in app_info_dict:
            app_info_dict["Default"] = self.init_empty_dict("OsApplication", "Os", "OsApplication(s)")

        # 解析APP_TASK
        app_task_sheet = self.workbook["APP_TASK"]
        app_task_header_column = {
            "OsApps": 1,
            "reference":
            {
                "OsTasks": 2
            }
        }
        temp_info_dict = self.get_sheet_info_by_header(app_task_sheet, app_task_header_column, "OsApps")

        # 优先以APP_TASK为准
        app_header_column = {
            "OsApplication(s)": 1,
            "parameter":
            {
                "OsIsPrivileged": 2,
                "OsAppEcucPartitionRef": 3,
                "OsApplicationCoreRef": 4,
            },
            "reference":
            {
                "OsAppCounterRef": 5,
                "OsAppTaskRef": 6,
                "OsAppScheduleTableRef": 7,
                "OsAppIsrRef": 8,
            }
        }
        header_mapping = {
            "OsTasks": "OsAppTaskRef",
        }
        for app_name in temp_info_dict.keys():
            if app_name not in app_info_dict:
                app_info_dict[app_name] = {}
                for k in temp_info_dict[app_name].keys():
                    app_info_dict[app_name][header_mapping[k]] = copy.deepcopy(temp_info_dict[app_name][k])
                # 用默认值补充缺少的key
                for k in app_header_column["parameter"].keys():
                    if k not in app_info_dict[app_name]:
                        app_info_dict[app_name][k] = copy.deepcopy(app_info_dict["Default"][k])
                for k in app_header_column["reference"].keys():
                    if k not in app_info_dict[app_name]:
                        app_info_dict[app_name][k] = copy.deepcopy(app_info_dict["Default"][k])
            else:
                # 覆盖特定key
                for k in temp_info_dict[app_name].keys():
                    app_info_dict[app_name][header_mapping[k]] = copy.deepcopy(temp_info_dict[app_name][k])

        for app_name in app_info_dict.keys():
            app_info_dict[app_name]["OsAppIsrRef"] += self.app_isr_mapping.get(app_name, [])
            app_info_dict[app_name]["OsAppIsrRef"] = OsCfgExecutor.remove_list_duplicates(app_info_dict[app_name]["OsAppIsrRef"])

        os_bswmd = self.autosar_bswmd_class["Os"]
        os_ecuc = self.autosar_bsw_class["Os"]
        self.create_os_app(os_ecuc, os_bswmd, app_info_dict)

    def create_os_app(self, os_ecuc, os_bswmd, app_info_dict):
        """配置OsApplication(s)下的app"""
        app_header_info = self.format_info[self.excel_version]["sheet"]["OsApplication"]["header"]
        # 移除原有app
        OsCfgExecutor.remove_containers(os_ecuc, "/AUTOSAR/EcucDefs/Os/OsApplication", True)
        os_container = os_ecuc.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0]
        # 创建app
        if "Default" in app_info_dict:
            del app_info_dict["Default"]
        OsCfgExecutor.add_containers(self.format_info, self.excel_version, "OsApplication", os_bswmd, app_info_dict, os_container)

class OsCfgExecutor:
    @classmethod
    def remove_list_duplicates(cls, lst):
        seen = set()
        new_lst = []
        for item in lst:
            if item not in seen:
                new_lst.append(item)
                seen.add(item)
        return new_lst

    @classmethod
    def find_container_by_def(cls, obj, def_str, parent_def_str=None):
        """获取所有指定container"""
        res = []
        if obj is None:
            return res
        parent_obj = None
        if parent_def_str is not None:
            for _ in obj:
                if _.DEFINITION_REF.valueOf_ == parent_def_str:
                    parent_obj = _.SUB_CONTAINERS.ECUC_CONTAINER_VALUE
                    break
        else:
            parent_obj = obj
        if parent_obj is None:
            return res

        for _ in parent_obj:
            if _.DEFINITION_REF.valueOf_ == def_str:
                res.append(_)
        return res

    @classmethod
    def find_value_by_def(cls, obj, def_str, mod_arxml):
        """获取所有指定参数值"""
        mod_obj = AutosarUtil.find(mod_arxml, def_str)
        tag = mod_obj.original_tagname_
        if tag in ECUC_PARAM_VALUE_LIST:
            param_obj = None
            if tag in ECUC_NUMERICAL_PARAM_VALUE_LIST:
                param_obj = obj.PARAMETER_VALUES.ECUC_NUMERICAL_PARAM_VALUE
            elif tag in ECUC_TEXTUAL_PARAM_VALUE_LIST:
                param_obj = obj.PARAMETER_VALUES.ECUC_TEXTUAL_PARAM_VALUE
            elif tag in ECUC_ADD_INFO_PARAM_VALUE_LIST:
                param_obj = obj.PARAMETER_VALUES.ECUC_ADD_INFO_PARAM_VALUE
            else:
                logger.warning(f"{tag} is not supported.")
                return ""
            for _ in param_obj:
                if _.DEFINITION_REF.valueOf_ == def_str:
                    if _.VALUE is not None:
                        return _.VALUE.valueOf_
        elif tag in ECUC_REF_VALUE_LIST:
            ref_list = []
            if tag in ECUC_REFERENCE_VALUE_LIST:
                for _ in obj.REFERENCE_VALUES.ECUC_REFERENCE_VALUE:
                    if _.DEFINITION_REF.valueOf_ == def_str:
                        if _.VALUE_REF is not None:
                            ref_list.append(_.VALUE_REF.valueOf_)
            elif tag in ECUC_IREFERENCE_VALUE_LIST:
                for _ in obj.REFERENCE_VALUES.ECUC_INSTANCE_REFERENCE_VALUE:
                    if _.DEFINITION_REF.valueOf_ == def_str:
                        if _.VALUE_IREF is not None:
                            ref_list.append(_.VALUE_IREF.TARGET_REF.valueOf_)
            else:
                logger.warning(f"{tag} is not supported.")
            ref_list = [_ for _ in ref_list if _ is not None]
            return ref_list
        else:
            logger.warning(f"{tag} is not supported.")
            return ""

    @classmethod
    def set_value(cls, obj, value, is_lower=False):
        if value is None:
            return
        if is_lower is True or str(value).lower() == "true" or str(value).lower() == "false":
            utils.set_value(obj, str(value).lower())
        else:
            utils.set_value(obj, str(value))

    @classmethod
    def remove_containers(cls, ecuc_obj, def_str, ignore_automatic=False):
        """移除已有container"""
        container_list = ecuc_obj.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0].CONTAINERS.ECUC_CONTAINER_VALUE
        if ignore_automatic is False:
            container_list = [item for item in container_list if item.DEFINITION_REF.valueOf_ != def_str]
        else:
            new_list = []
            for item in container_list:
                if item.DEFINITION_REF.valueOf_ != def_str:
                    new_list.append(item)
                    continue
                if item.ANNOTATIONS is None:
                    continue
                for _ in item.ANNOTATIONS.ANNOTATION:
                    if _.ANNOTATION_ORIGIN.valueOf_ == "IsAutomatic":
                        new_list.append(item)
                        break
            container_list = new_list
        ecuc_obj.AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0].CONTAINERS.ECUC_CONTAINER_VALUE = container_list

    @classmethod
    def add_reference_list(cls, item_list, parent, def_str, suffix):
        """添加多个reference配置"""
        if isinstance(item_list, list) is False or len(item_list) == 0:
            return
        item_list = OsCfgExecutor.remove_list_duplicates(item_list)
        ref_list = parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE
        for item in item_list:
            # 不能重复
            is_existed = False
            for ref in ref_list:
                if ref.DEFINITION_REF.valueOf_ != def_str:
                    continue
                if ref.VALUE_REF is None:
                    continue
                if ref.VALUE_REF.valueOf_ == str(item):
                    is_existed = True
                    break
            if is_existed is False:
                new_ref = autosar44.ECUC_REFERENCE_VALUE(
                    DEFINITION_REF=autosar44.DEFINITION_REFType2135(valueOf_=def_str)
                )
                new_ref.original_tagname_ = "ECUC-REFERENCE-VALUE"
                item = str(item)
                if "/" in item or "<empty>" in item:
                    utils.set_value_ref(new_ref, item)
                else:
                    utils.set_value_ref(new_ref, f"{suffix}{item}")
                ref_list.append(new_ref)
        parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE = ref_list

    @classmethod
    def modify_reference_list(cls, item_list, parent, def_str, suffix):
        """修改多个reference配置"""
        if isinstance(item_list, str) is True:
            item_list = [item_list]
        if isinstance(item_list, list) is False or len(item_list) == 0:
            return
        new_list = []
        for item in item_list:
            item = str(item)
            if "/" in item or "<empty>" in item:
                new_list.append(item)
            else:
                new_list.append(f"{suffix}{item}")
        item_list = new_list
        item_list = OsCfgExecutor.remove_list_duplicates(item_list)
        parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE = [
            item for item in parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE if \
                item.DEFINITION_REF.valueOf_ != def_str
        ]
        for item in item_list:
            ref = autosar44.ECUC_REFERENCE_VALUE(
                DEFINITION_REF=autosar44.DEFINITION_REFType2135(valueOf_=def_str)
            )
            ref.original_tagname_ = "ECUC-REFERENCE-VALUE"
            utils.set_value_ref(ref, str(item))
            parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE.append(ref)

    @classmethod
    def modify_single_reference(cls, item, parent, def_str, suffix):
        """修改单个reference配置"""
        if item is None:
            return
        item = str(item)
        for ref in parent.REFERENCE_VALUES.ECUC_REFERENCE_VALUE:
            if ref.DEFINITION_REF.valueOf_ == def_str:
                if "/" in item or "<empty>" in item:
                    utils.set_value_ref(ref, item)
                else:
                    utils.set_value_ref(ref, f"{suffix}{item}")
                break

    @classmethod
    def add_containers(cls, format_info, excel_version, sheet_name, bswmd_obj, info_dict, parent_container):
        """修改配置"""
        header = format_info[excel_version]["sheet"][sheet_name]["header"]
        # 参数类型
        param_tag_info = {}
        for param_name, item in header.items():
            param_mod_path = item[1]
            if param_mod_path is not None:
                param_tag_info[param_name] = bswmd_obj.find(param_mod_path).original_tagname_

        config_releation = format_info[excel_version]["sheet"][sheet_name]["config_releation"]
        for container_name, param_list in config_releation.items():
            container_mod_path = header[container_name][1]
            container_bswmd_obj = bswmd_obj.find(container_mod_path)
            for name, param_info in info_dict.items():
                # 创建container
                container = AutosarUtil.create_container_value(str(name), container_bswmd_obj, parent_container, str(name))
                # 修改配置值
                for param_name in param_list:
                    param_mod_path = header[param_name][1]
                    if param_mod_path is None:
                        continue
                    value = param_info[param_name]
                    tag = param_tag_info[param_name]
                    if tag in ECUC_PARAM_VALUE_LIST:
                        param_obj = None
                        if tag in ECUC_NUMERICAL_PARAM_VALUE_LIST:
                            param_obj = container.PARAMETER_VALUES.ECUC_NUMERICAL_PARAM_VALUE
                        elif tag in ECUC_TEXTUAL_PARAM_VALUE_LIST:
                            param_obj = container.PARAMETER_VALUES.ECUC_TEXTUAL_PARAM_VALUE
                        elif tag in ECUC_ADD_INFO_PARAM_VALUE_LIST:
                            param_obj = container.PARAMETER_VALUES.ECUC_ADD_INFO_PARAM_VALUE
                        else:
                            logger.warning(f"{tag} is not supported.")
                            continue
                        for _ in param_obj:
                            if _.DEFINITION_REF.valueOf_ == param_mod_path:
                                cls.set_value(_, value)
                    elif tag in ECUC_REF_VALUE_LIST:
                        cls.modify_reference_list(value, container, param_mod_path, header[param_name][2])
                    else:
                        logger.warning(f"{tag} is not supported.")
                        continue


class OsCfgExcelExporter:
    def __init__(self, data_management) -> None:
        self.data_management = data_management
        self.autosar_bsw_class = data_management.autosar_bsw_class
        self.autosar_bswmd_class = data_management.autosar_bswmd_class

        self.format_info = {}           # excel格式信息
        self.workbook = None            # excel文件
        self.os_excel_path = None       # excel文件路径
        self.excel_version = None       # excel文件版本号

        self.alarm_app_mapping = {}     # alarm所属application
        self.isr_app_mapping = {}       # isr所属application
        self.event_task_mapping = {}    # event所属Task

    def export_excel(self, os_excel_path) -> bool:
        """导出OS现有配置"""
        self.os_excel_path = os_excel_path
        self.workbook = Workbook()

        if self.validate() is False:
            return False
        self.parse_config()

        try:
            self.workbook.save(self.os_excel_path)
        except PermissionError:
            logger.ui_error(f"{self.os_excel_path} is open in another application. Please close it and try again.")
            return False
        return True

    def validate(self) -> bool:
        """校验能否导出excel"""
        # 导出路径不能为空
        if not self.os_excel_path:
            logger.ui_error("File path can not be empty.")
            return False
        # 导出的excel格式
        excel_format_path = os.path.join(CUR_FILE_PATH, "excel_format.json")
        if not os.path.exists(excel_format_path):
            logger.ui_error(f"'{excel_format_path}' does not exist.")
            return False
        with open(excel_format_path, "r", encoding="utf-8") as f:
            self.format_info = json.load(f)
        # 导出的excel版本号
        for excel_version, item in self.format_info.items():
            if item.get("is_exported", False) is True:
                self.excel_version = excel_version
                break
        if self.excel_version is None:
            logger.ui_error(f"There is at least and only one version in VERSION-INFO where 'is_exported' is True.")
            return False
        # 导出的模块
        for module in self.format_info[self.excel_version]["releated_module_list"]:
            if module not in self.autosar_bswmd_class:
                logger.ui_error(f"{module} module exists in 'releated_module_list',but {module}_bswmd.arxml does not exist in current project.")
                return False
        # 表头的mod arxml格式
        for sheet_name, sheet_info in self.format_info[self.excel_version]["sheet"].items():
            for header, item in sheet_info["header"].items():
                mod_path = item[1]
                if mod_path is None:
                    continue
                module = mod_path.split("/")[3]
                if module not in self.autosar_bswmd_class:
                    logger.ui_error(f"{module}_bswmd.arxml does not exist in current project.")
                    return False
                if AutosarUtil.find(self.autosar_bswmd_class[module], mod_path) is None:
                    logger.ui_error(f"{header}'s mod_path {mod_path} does not exist in {module}_bswmd.arxml.")
                    return False
        return True

    def parse_config(self):
        """解析工程现有配置并存入excel"""
        self.parse_version()
        self.parse_ecuc_core()
        self.parse_ecuc_partition()
        self.parse_app_task()
        self.parse_isr()
        self.parse_event()
        if "Sheet" in self.workbook.sheetnames:
            self.workbook.remove(self.workbook["Sheet"])

    def parse_sheet(self, sheet_name, module_name) -> dict:
        """结构化解析sheet信息"""
        sheet_info = {}
        if module_name not in self.autosar_bsw_class:
            return sheet_info

        header = self.format_info[self.excel_version]["sheet"][sheet_name]["header"]
        config_releation = self.format_info[self.excel_version]["sheet"][sheet_name]["config_releation"]
        module_obj = self.autosar_bsw_class[module_name].AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0].CONTAINERS.ECUC_CONTAINER_VALUE
        for container_name, param_list in config_releation.items():
            container_mod_path = header[container_name][1]
            parent_mod_path = None
            if len(container_mod_path.split("/")) >= 6:
                parent_mod_path = os.path.dirname(container_mod_path)
            container_list = OsCfgExecutor.find_container_by_def(module_obj, container_mod_path, parent_mod_path)
            for container in container_list:
                if AutosarUtil.is_automatic(container) is True:
                    continue
                sheet_info[container.SHORT_NAME.valueOf_] = {}
                for param in param_list:
                    value = OsCfgExecutor.find_value_by_def(container, header[param][1], self.autosar_bswmd_class[module_name])
                    sheet_info[container.SHORT_NAME.valueOf_][param] = value
        return sheet_info

    def construct_sheet(self, sheet_name, sheet_info, name_key):
        """构造sheet"""
        if sheet_name in self.workbook.sheetnames:
            self.workbook.remove(self.workbook[sheet_name])
        sheet = self.workbook.create_sheet(title=sheet_name)

        header_column = {}
        last_row = -1
        for header_name, item in self.format_info[self.excel_version]["sheet"][sheet_name]["header"].items():
            ExcelExecutor.write_sheet_header(sheet, item[0], header_name)
            header_column[header_name] = item[0][2]
            last_row = max(last_row, item[0][0])
        for name, item in sheet_info.items():
            fill = None
            if name == "Default":
                fill = PatternFill(start_color="FF00D6B9", end_color="FF00D6B9", fill_type="solid")
            ExcelExecutor.write_sheet_cell(sheet, last_row + 1, header_column[name_key], name, fill)
            cnt = 1
            for k, v in item.items():
                if isinstance(v, list) is True:
                    for i in range(len(v)):
                        ExcelExecutor.write_sheet_cell(sheet, last_row + 1 + i, header_column[k], v[i], fill)
                    cnt = max(cnt, len(v))
                else:
                    ExcelExecutor.write_sheet_cell(sheet, last_row + 1, header_column[k], v, fill)
                    cnt = max(cnt, 1)
            sheet.merge_cells(start_row=last_row + 1, end_row=last_row + cnt, start_column=header_column[name_key], end_column=header_column[name_key])
            for k, v in item.items():
                if isinstance(v, list) is False:
                    sheet.merge_cells(start_row=last_row + 1, end_row=last_row + cnt, start_column=header_column[k], end_column=header_column[k])
            last_row += cnt
        ExcelExecutor.set_column_widths(sheet)
        ExcelExecutor.set_row_heights(sheet)

        return sheet

    def parse_version(self):
        """版本"""
        sheet_name = "Version"
        if sheet_name in self.workbook.sheetnames:
            self.workbook.remove(self.workbook[sheet_name])
        sheet = self.workbook.create_sheet(title=sheet_name)
        ExcelExecutor.write_sheet_header(sheet, (1, 1, 1, 1), "Version")
        ExcelExecutor.write_sheet_cell(sheet, 2, 1, self.excel_version)
        ExcelExecutor.set_column_widths(sheet)
        ExcelExecutor.set_row_heights(sheet)

    def parse_ecuc_core(self):
        """解析工程现有EcucCoreDefinition配置"""
        sheet_info = self.parse_sheet("EcucCoreDefinition", "EcuC")
        self.construct_sheet("EcucCoreDefinition", sheet_info, "EcucCoreDefinition(s)")

    def parse_ecuc_partition(self):
        """解析工程现有EcucPartition配置"""
        sheet_info = self.parse_sheet("EcucPartition", "EcuC")
        self.construct_sheet("EcucPartition", sheet_info, "EcucPartition(s)")

    def parse_app_task(self):
        """解析工程现有application和task配置"""
        if "Os" in self.autosar_bsw_class:
            os_obj = self.autosar_bsw_class["Os"].AR_PACKAGES.AR_PACKAGE[0].ELEMENTS.ECUC_MODULE_CONFIGURATION_VALUES[0].CONTAINERS.ECUC_CONTAINER_VALUE
        else:
            os_obj = None
        sheet_info = self.format_info[self.excel_version]["sheet"]
        # OsApplication
        app_sheet_info = self.parse_sheet("OsApplication", "Os")
        # app下的ISR和alarm
        self.alarm_app_mapping = {}
        self.isr_app_mapping = {}
        for app, item in app_sheet_info.items():
            for isr in item["OsAppIsrRef"]:
                if not isr:
                    continue
                self.isr_app_mapping[isr.split("/")[-1]] = app
        # OsTask
        task_sheet_info = self.parse_sheet("OsTask", "Os")
        task_mod_path = sheet_info["OsTask"]["header"]["OsTask(s)"][1]
        task_list = OsCfgExecutor.find_container_by_def(os_obj, task_mod_path)
        autostart_mod_path = sheet_info["OsTask"]["header"]["OsTaskAutostart"][1]
        for task in task_list:
            for sub in task.SUB_CONTAINERS.ECUC_CONTAINER_VALUE:
                if sub.DEFINITION_REF.valueOf_ == autostart_mod_path:
                    for ref in sub.REFERENCE_VALUES.ECUC_REFERENCE_VALUE:
                        if ref.DEFINITION_REF.valueOf_ == f"{autostart_mod_path}/OsTaskAppModeRef":
                            task_sheet_info[task.SHORT_NAME.valueOf_]["OsTaskAutostart"] = ref.VALUE_REF.valueOf_
        self.event_task_mapping = {}
        for task, item in task_sheet_info.items():
            for event in item["OsTaskEventRef"]:
                if not event:
                    continue
                self.event_task_mapping[event.split("/")[-1]] = task
        # APP_TASK
        mapping = {
            "Task Type": "OsTaskType",
            "Task Priority": "OsTaskPriority",
            "Task Stack Size [Byte]": "OsTaskStackSize",
        }
        app_task_sheet_info = {}
        for app, app_info in app_sheet_info.items():
            app_task_sheet_info[app] = {
                "OsTasks": [],
                "Task Event": [],
                "Task所属": []
            }
            for key in mapping:
                app_task_sheet_info[app][key] = []
            for task in app_info["OsAppTaskRef"]:
                task_name = task.split("/")[-1]
                app_task_sheet_info[app]["OsTasks"].append(task_name)
                app_task_sheet_info[app]["Task Event"].append("、\n".join(task_sheet_info[task_name]["OsTaskEventRef"]))
                app_task_sheet_info[app]["Task所属"].append(None)
                for key in mapping:
                    app_task_sheet_info[app][key].append(task_sheet_info[task_name][mapping[key]])
        self.construct_sheet("APP_TASK", app_task_sheet_info, "OsApps")
        default = {
            "Default":
            {
                "OsIsPrivileged": "true",
                "OsAppEcucPartitionRef": None,
                "OsApplicationCoreRef": None,
                "OsAppCounterRef": [""],
                "OsAppTaskRef": [""],
                "OsAppScheduleTableRef": [""],
                "OsAppIsrRef": [""],
            }
        }
        app_sheet_info = {**default, **app_sheet_info}
        self.construct_sheet("OsApplication", app_sheet_info, "OsApplication(s)")
        default = {
            "Default":
            {
                "OsTaskUsesFpu": "true",
                "OsTaskSchedule": "FULL",
                "OsTaskType": "BASIC",
                "OsTaskActivation": None,
                "OsTaskPriority": None,
                "OsTaskStackSize": None,
                "OsTaskEventRef": [""],
                "OsTaskAccessingApplication": [""],
                "OsTaskAutostart": None
            }
        }
        task_sheet_info = {**default, **task_sheet_info}
        self.construct_sheet("OsTask", task_sheet_info, "OsTask(s)")

    def parse_isr(self):
        """解析工程现有isr配置"""
        sheet_info = self.parse_sheet("OsIsr", "Os")
        for isr in sheet_info:
            sheet_info[isr]["所属Application"] = self.isr_app_mapping.get(isr, None)
        self.construct_sheet("OsIsr", sheet_info, "OsIsr(s)")

    def parse_event(self):
        """解析工程现有event配置"""
        sheet_info = self.parse_sheet("OsEvent", "Os")
        for event in sheet_info:
            sheet_info[event]["所属Task"] = self.event_task_mapping.get(event, None)
        self.construct_sheet("OsEvent", sheet_info, "OsEvent(s)")


class ExcelExecutor:
    """操作excel文件"""
    @classmethod
    def write_sheet_cell(cls, sheet, row, column, value, fill=None, font=Font(), border=None):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font
        if border is not None:
            cell.border = border

    @classmethod
    def write_sheet_header(cls, sheet, coor, value):
        fill = PatternFill(start_color="14C0FF", end_color="14C0FF", fill_type="solid")
        font = Font(name="Cascadia Mono", size=12, bold=True)
        border = Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000")
        )
        start_row, end_row, start_column, end_column = coor[0], coor[1], coor[2], coor[3]
        cell = cls.write_sheet_cell(sheet, start_row, start_column, value, fill, font, border)
        sheet.merge_cells(start_row=start_row, end_row=end_row, start_column=start_column, end_column=end_column)

    @classmethod
    def set_row_heights(cls, sheet):
        for row in sheet.rows:
            row_height = 18
            for cell in row:
                try:
                    lines = str(cell.value).count("\n") + 1
                    adjusted_height = lines * 15
                    if adjusted_height > row_height:
                        row_height = adjusted_height
                except:
                    pass
            sheet.row_dimensions[row[0].row].height = row_height

    @classmethod
    def set_column_widths(cls, sheet, ori_length=0):
        for column in sheet.columns:
            max_length = ori_length
            for _ in column:
                if not hasattr(_, "column_letter"):
                    continue
                column_name = _.column_letter
                break
            for cell in column:
                for _ in str(cell.value).split("\n"):
                    if len(_) > max_length:
                        max_length = len(_)
            adjusted_width = (max_length + 2) * 1.2
            sheet.column_dimensions[column_name].width = adjusted_width

